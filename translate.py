import json, os, sys, time, requests, re
from docx import Document

CHECKPOINT_FILE = "progress.json"
OPENROUTER_KEY = "sk-or-..."
MODEL = "google/gemini-2.5-pro"

def load_progress():
    if os.path.exists(CHECKPOINT_FILE):
        with open(CHECKPOINT_FILE, "r") as f:
            return json.load(f)
    return {}

def save_progress(progress):
    with open(CHECKPOINT_FILE, "w") as f:
        json.dump(progress, f, ensure_ascii=False, indent=2)

def paragraph_to_tagged(paragraph):
    return "".join(
        f'<r id="{i}">{run.text}</r>'
        for i, run in enumerate(paragraph.runs)
        if run.text
    )

def apply_tagged_translation(paragraph, tagged_translation):
    matches = re.findall(r'<r id="(\d+)">(.*?)</r>', tagged_translation, re.DOTALL)
    translated_map = {int(idx): text for idx, text in matches}
    for i, run in enumerate(paragraph.runs):
        if i in translated_map:
            run.text = translated_map[i]

def translate_with_retry(tagged_text, context, retries=3):
    for attempt in range(retries):
        try:
            response = requests.post(
                "https://openrouter.ai/api/v1/chat/completions",
                headers={"Authorization": f"Bearer {OPENROUTER_KEY}"},
                json={
                    "model": MODEL,
                    "messages": [
                        {"role": "system", "content": (
                            "Traduci EN→IT mantenendo tono e registro. "
                            "Il testo è suddiviso in porzioni marcate con tag <r id=\"N\">...</r>. "
                            "Mantieni tutti i tag esattamente invariati; traduci solo il testo al loro interno. "
                            "Restituisci solo il testo con i tag intatti, senza aggiungere nulla fuori dai tag."
                        )},
                        {"role": "user", "content": f"CONTESTO:\n{context}\n\nTESTO:\n{tagged_text}"}
                    ],
                    "temperature": 0.2
                },
                timeout=60
            )
            return response.json()["choices"][0]["message"]["content"].strip()
        except Exception as e:
            print(f"  ⚠️ Tentativo {attempt+1} fallito: {e}")
            time.sleep(5 * (attempt + 1))
    raise Exception(f"Paragrafo fallito dopo {retries} tentativi")

def collect_paragraphs(doc):
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    return paragraphs

def translate_docx(input_path, output_path):
    doc = Document(input_path)
    all_paragraphs = collect_paragraphs(doc)
    progress = load_progress()

    for key, tagged_translation in progress.items():
        idx = int(key)
        if idx < len(all_paragraphs):
            apply_tagged_translation(all_paragraphs[idx], tagged_translation)

    total = sum(
        1 for i, p in enumerate(all_paragraphs)
        if p.text.strip() and str(i) not in progress
    )
    done = 0

    for para_idx, paragraph in enumerate(all_paragraphs):
        key = str(para_idx)
        if not paragraph.text.strip() or key in progress:
            continue

        tagged = paragraph_to_tagged(paragraph)
        if not tagged:
            continue

        context = "\n".join(
            all_paragraphs[j].text.strip()
            for j in range(max(0, para_idx - 2), para_idx)
            if all_paragraphs[j].text.strip()
        )

        done += 1
        print(f"  🔄 Paragrafo {done}/{total}...")

        result = translate_with_retry(tagged, context)
        apply_tagged_translation(paragraph, result)

        progress[key] = result
        save_progress(progress)
        time.sleep(0.8)

    doc.save(output_path)
    print("\n✅ Traduzione completata!")

if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Uso: python translate.py input.docx output.docx")
        sys.exit(1)
    translate_docx(sys.argv[1], sys.argv[2])
